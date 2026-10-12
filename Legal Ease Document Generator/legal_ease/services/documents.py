"""DOCX and PDF export helpers for generated or user-edited drafts."""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF
from fpdf.enums import XPos, YPos


def _lines(text: str) -> list[str]:
    return [line.strip() for line in text.replace("\r\n", "\n").split("\n")]


def create_docx(text: str, document_type: str = "Legal document") -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(document_type.strip() or "Legal document")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(31, 56, 51)

    for line in _lines(text):
        if not line:
            continue
        if line.startswith("## ") or (line[:1].isdigit() and ". " in line[:5]):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line.lstrip("# "))
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase | Draft for review by qualified counsel").font.size = Pt(8)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _LegalEasePDF(FPDF):
    def header(self) -> None:
        self.set_font("Helvetica", "B", 9)
        self.set_text_color(35, 80, 70)
        self.cell(0, 7, "LEGALEASE  |  DOCUMENT DRAFT", align="C")
        self.ln(11)
        self.set_text_color(30, 36, 34)

    def footer(self) -> None:
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(90, 100, 96)
        self.cell(0, 8, f"Draft for legal review  |  Page {self.page_no()}", align="C")


def create_pdf(text: str, document_type: str = "Legal document") -> bytes:
    pdf = _LegalEasePDF()
    pdf.set_margins(22, 20, 22)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(
        0,
        9,
        _pdf_safe(document_type.strip() or "Legal document"),
        align="C",
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )
    pdf.ln(6)

    for line in _lines(text):
        if not line:
            pdf.ln(3)
        elif line.startswith("## ") or (line[:1].isdigit() and ". " in line[:5]):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(
                0,
                6,
                _pdf_safe(line.lstrip("# ")),
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )
        else:
            is_bullet = line.startswith(("- ", "* "))
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(
                0,
                5.5,
                _pdf_safe(("- " + line[2:]) if is_bullet else line),
                new_x=XPos.LMARGIN,
                new_y=YPos.NEXT,
            )

    return bytes(pdf.output())


def _pdf_safe(value: str) -> str:
    return value.encode("latin-1", errors="replace").decode("latin-1")
