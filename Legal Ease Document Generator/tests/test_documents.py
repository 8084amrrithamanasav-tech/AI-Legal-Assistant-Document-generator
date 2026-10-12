from io import BytesIO

from docx import Document

from legal_ease.services.documents import create_docx, create_pdf


def test_docx_contains_title_and_document_text() -> None:
    content = create_docx("## 1. Payment\nPayment is due within 30 days.", "Services Agreement")
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert "Services Agreement" in paragraphs
    assert "1. Payment" in paragraphs
    assert "Payment is due within 30 days." in paragraphs
    assert "Draft for review by qualified counsel" in document.sections[0].footer.paragraphs[0].text


def test_pdf_is_generated_with_pdf_signature() -> None:
    content = create_pdf("## 1. Payment\nPayment is due within 30 days.", "Services Agreement")

    assert content.startswith(b"%PDF-")
